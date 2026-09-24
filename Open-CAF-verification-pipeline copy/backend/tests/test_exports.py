"""Tests for Multi-Format Exporters: CSV (RFC 4180), Excel (.xlsx), Jira REST JSON, and GitHub Issues."""
import pytest
import io
import csv
from datetime import date, timedelta
import openpyxl
from httpx import AsyncClient, ASGITransport
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker, AsyncSession
from sqlalchemy.orm import Session

from app.main import app
from app.core.database import Base, get_db
from app.core.security import get_password_hash, create_access_token
from app.models.tenant import CouncilTenant, AuthorityType
from app.models.user import User, UserRole
from app.models.remediation import RemediationTask, RemediationStatus, RemediationPriority
from app.schemas.assessment import AssessmentCreate
from app.services.assessment_service import AssessmentService
from app.services.export_service import ExportService
from app.cli.seed_caf import seed_caf_data

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"


@pytest.fixture(scope="function")
async def export_test_env():
    """Initializes in-memory test DB, seeds CAF v4.0, creates assessment, users, and tasks."""
    engine = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    # Seed CAF Taxonomy
    async with engine.connect() as conn:
        def do_seed_caf(connection):
            sync_sess = Session(bind=connection)
            seed_caf_data(session=sync_sess)
            sync_sess.close()
        await conn.run_sync(do_seed_caf)
        await conn.commit()

    async_session = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)

    async with async_session() as session:
        # Tenant A: Borsetshire Council
        tenant_a = CouncilTenant(name="Borsetshire Council", authority_type=AuthorityType.UNITARY)
        session.add(tenant_a)

        # Tenant B: Barsetshire County
        tenant_b = CouncilTenant(name="Barsetshire County", authority_type=AuthorityType.COUNTY)
        session.add(tenant_b)
        await session.flush()

        # Users
        ciso_a = User(
            tenant_id=tenant_a.id,
            email="ciso@borsetshire.gov.uk",
            hashed_password=get_password_hash("Pass123!"),
            full_name="Arthur CISO",
            role=UserRole.CISO_ADMIN,
            is_active=True,
        )
        ciso_b = User(
            tenant_id=tenant_b.id,
            email="ciso@barsetshire.gov.uk",
            hashed_password=get_password_hash("Pass123!"),
            full_name="Boris CISO",
            role=UserRole.CISO_ADMIN,
            is_active=True,
        )
        session.add_all([ciso_a, ciso_b])
        await session.flush()

        # Create assessment in Tenant A
        ass = await AssessmentService.create_assessment(
            db=session,
            tenant_id=tenant_a.id,
            payload=AssessmentCreate(
                title="Borsetshire Annual Resilience Audit 2026",
                scope_description="All corporate and statutory service networks across Borsetshire Council",
                council_service_name="Revenues and Corporate IT",
            ),
        )

        # Seed remediation tasks
        task1 = RemediationTask(
            tenant_id=tenant_a.id,
            assessment_id=ass.id,
            outcome_id="B2.a",
            title='Enforce FIDO2 MFA on "Legacy Gateway", VPN & RDP portals',
            description="Replace single-factor admin logins with Entra ID conditional access and hardware tokens.",
            status=RemediationStatus.IN_PROGRESS,
            priority=RemediationPriority.CRITICAL,
            assigned_owner_name="David Cameron",
            assigned_owner_email="david.cameron@borsetshire.gov.uk",
            estimated_effort_hours=45.0,
            estimated_cost_gbp=18500.0,
            target_completion_date=date.today() + timedelta(days=14),
            external_ticket_id="JIRA-2041",
            technical_steps=[
                {"step": "Configure Conditional Access", "completed": True},
                {"step": "Distribute YubiKeys", "completed": False},
            ],
        )

        task2 = RemediationTask(
            tenant_id=tenant_a.id,
            assessment_id=ass.id,
            outcome_id="B4.a",
            title="Deploy Air-Gapped Immutable Backups for Revenues DB",
            description="Implement S3 object lock and weekly restoration testing sandbox.",
            status=RemediationStatus.BACKLOG,
            priority=RemediationPriority.CRITICAL,
            assigned_owner_name="Sarah Jenkins",
            assigned_owner_email="sarah.jenkins@borsetshire.gov.uk",
            estimated_effort_hours=60.0,
            estimated_cost_gbp=32000.0,
            target_completion_date=date.today() + timedelta(days=28),
            external_ticket_id="JIRA-2049",
            technical_steps=[
                {"step": "Procure storage tier", "completed": True},
            ],
        )

        task3 = RemediationTask(
            tenant_id=tenant_a.id,
            assessment_id=ass.id,
            outcome_id="D1.a",
            title="Conduct Ransomware Crisis Tabletop Drill with Cabinet",
            description="Exercise incident response plan D1-v3 with Gold Command.",
            status=RemediationStatus.COMPLETED,
            priority=RemediationPriority.HIGH,
            assigned_owner_name="Rachel Sterling",
            assigned_owner_email="rachel.sterling@borsetshire.gov.uk",
            estimated_effort_hours=25.0,
            estimated_cost_gbp=6500.0,
            target_completion_date=date.today() - timedelta(days=5),
            external_ticket_id="CAB-102",
        )

        session.add_all([task1, task2, task3])
        await session.commit()

        token_a = create_access_token(ciso_a.id, tenant_a.id, ciso_a.role.value)
        token_b = create_access_token(ciso_b.id, tenant_b.id, ciso_b.role.value)

        env = {
            "engine": engine,
            "session_factory": async_session,
            "tenant_a_id": tenant_a.id,
            "tenant_b_id": tenant_b.id,
            "assessment_id": ass.id,
            "token_a": token_a,
            "token_b": token_b,
        }

    async def override_get_db():
        async with async_session() as s:
            yield s

    app.dependency_overrides[get_db] = override_get_db
    yield env
    app.dependency_overrides.clear()
    await engine.dispose()


@pytest.mark.asyncio
async def test_export_remediation_csv_endpoint(export_test_env):
    """Verifies RFC 4180 compliant CSV export format, headers, and quote escaping."""
    env = export_test_env
    headers = {"Authorization": f"Bearer {env['token_a']}"}
    ass_id = env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        resp = await ac.get(f"/api/v1/assessments/{ass_id}/export/csv", headers=headers)
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        assert f'attachment; filename="opencaf_remediation_{ass_id}.csv"' in resp.headers["content-disposition"]

        # Parse CSV content
        csv_reader = csv.reader(io.StringIO(resp.text))
        rows = list(csv_reader)

        # Assert headers
        assert len(rows) == 4  # Header + 3 data rows
        header_row = rows[0]
        assert "Task ID" in header_row
        assert "Contributing Outcome" in header_row
        assert "Title" in header_row
        assert "Status" in header_row
        assert "Priority" in header_row
        assert "Estimated Cost (GBP)" in header_row
        assert "Estimated Effort (Hours)" in header_row

        # Check data row with quotes and commas
        outcomes = [r[2] for r in rows[1:]]
        assert "B2.a" in outcomes
        assert "B4.a" in outcomes
        assert "D1.a" in outcomes

        # Verify task with quotes inside title was properly parsed
        row_b2a = next(r for r in rows[1:] if r[2] == "B2.a")
        assert 'Enforce FIDO2 MFA on "Legacy Gateway", VPN & RDP portals' == row_b2a[3]
        assert row_b2a[5] == "IN_PROGRESS"
        assert row_b2a[6] == "CRITICAL"
        assert float(row_b2a[10]) == 18500.0


@pytest.mark.asyncio
async def test_export_excel_endpoint_structure_and_styling(export_test_env):
    """Verifies openpyxl generated .xlsx file structure, titles, styles, and formula calculations."""
    env = export_test_env
    headers = {"Authorization": f"Bearer {env['token_a']}"}
    ass_id = env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        resp = await ac.get(f"/api/v1/assessments/{ass_id}/export/excel", headers=headers)
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert f"opencaf_remediation_action_plan_{ass_id}.xlsx" in resp.headers["content-disposition"]

        # Load workbook into openpyxl from binary stream
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "Remediation Action Plan" in wb.sheetnames
        ws = wb["Remediation Action Plan"]

        # Assert Title cell
        assert "Borsetshire Council" in ws["A1"].value
        assert "Open CAF" in ws["A1"].value

        # Assert Table Headers (Row 4)
        expected_headers = [
            "Ref #", "Outcome", "Remediation Action Item", "Priority", "Status",
            "Assigned Lead", "Lead Email", "Effort (Hrs)", "Est. Budget (£)",
            "Target Date", "Ticket ID"
        ]
        actual_headers = [ws.cell(row=4, column=c).value for c in range(1, 12)]
        assert actual_headers == expected_headers

        # Data rows (rows 5, 6, 7)
        outcome_codes = [ws.cell(row=r, column=2).value for r in (5, 6, 7)]
        assert "B2.a" in outcome_codes
        assert "B4.a" in outcome_codes
        assert "D1.a" in outcome_codes

        # Check total row formula (Row 8)
        assert ws.cell(row=8, column=1).value == "Total Action Plan Commitment:"
        assert "=SUM(H5:H7)" in str(ws.cell(row=8, column=8).value)
        assert "=SUM(I5:I7)" in str(ws.cell(row=8, column=9).value)

        # Check number format for currency on any data row
        assert "£" in ws.cell(row=5, column=9).number_format


@pytest.mark.asyncio
async def test_export_jira_payload_endpoint(export_test_env):
    """Verifies Jira Cloud/Server bulk format: issueUpdates, labels, priority mapping, and project keys."""
    env = export_test_env
    headers = {"Authorization": f"Bearer {env['token_a']}"}
    ass_id = env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # Export with custom project key and issue type
        resp = await ac.post(
            f"/api/v1/assessments/{ass_id}/export/jira-json",
            json={"project_key": "SEC", "issue_type": "SecurityTask"},
            headers=headers,
        )
        assert resp.status_code == 200
        data = resp.json()

        assert data["total_tasks"] == 3
        assert data["project_key"] == "SEC"
        assert "issueUpdates" in data
        issues = data["issueUpdates"]
        assert len(issues) == 3

        # Validate issue structure for B2.a task
        b2a_issue = next(i["fields"] for i in issues if "[B2.a]" in i["fields"]["summary"])
        assert b2a_issue["project"]["key"] == "SEC"
        assert b2a_issue["issuetype"]["name"] == "SecurityTask"
        assert "[B2.a]" in b2a_issue["summary"]
        assert b2a_issue["priority"]["name"] == "Highest"  # CRITICAL -> Highest
        assert "OpenCAF" in b2a_issue["labels"]
        assert "NCSC-CAF" in b2a_issue["labels"]
        assert "*NCSC CAF Contributing Outcome:* B2.a" in b2a_issue["description"]
        assert "*Estimated Budget:* £18,500.00" in b2a_issue["description"]


@pytest.mark.asyncio
async def test_export_github_payload_endpoint(export_test_env):
    """Verifies GitHub Issues batch payload format with tags, assignees, and checklist body."""
    env = export_test_env
    headers = {"Authorization": f"Bearer {env['token_a']}"}
    ass_id = env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        resp = await ac.post(f"/api/v1/assessments/{ass_id}/export/github-json", headers=headers)
        assert resp.status_code == 200
        data = resp.json()

        assert data["total_tasks"] == 3
        assert len(data["issues"]) == 3

        b2a_item = next(i for i in data["issues"] if "[B2.a]" in i["title"])
        assert "[B2.a]" in b2a_item["title"]
        assert "opencaf" in b2a_item["labels"]
        assert "outcome:b2.a" in b2a_item["labels"]
        assert "priority:critical" in b2a_item["labels"]
        assert "david.cameron" in b2a_item["assignees"]
        assert "Technical Checklist" in b2a_item["body"]


@pytest.mark.asyncio
async def test_export_tenant_isolation(export_test_env):
    """Verifies that Council Tenant B cannot export assessments belonging to Tenant A."""
    env = export_test_env
    headers_b = {"Authorization": f"Bearer {env['token_b']}"}
    ass_id_a = env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # CSV export cross-tenant
        resp_csv = await ac.get(f"/api/v1/assessments/{ass_id_a}/export/csv", headers=headers_b)
        assert resp_csv.status_code == 404

        # Excel export cross-tenant
        resp_xlsx = await ac.get(f"/api/v1/assessments/{ass_id_a}/export/excel", headers=headers_b)
        assert resp_xlsx.status_code == 404

        # Jira export cross-tenant
        resp_jira = await ac.post(f"/api/v1/assessments/{ass_id_a}/export/jira-json", headers=headers_b)
        assert resp_jira.status_code == 404


@pytest.mark.asyncio
async def test_export_unauthenticated_rejected(export_test_env):
    """Verifies that unauthenticated export requests are rejected with 401 Unauthorized."""
    ass_id = export_test_env["assessment_id"]

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        resp = await ac.get(f"/api/v1/assessments/{ass_id}/export/csv")
        assert resp.status_code == 401
